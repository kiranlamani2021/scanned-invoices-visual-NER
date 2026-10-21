import io
import tempfile
import unittest
from pathlib import Path

import cv2
import numpy as np
from PIL import Image

import app as invoice_app


class InvoiceAppTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.original_database_path = invoice_app.DATABASE_PATH
        self.original_upload_folder = invoice_app.app.config["UPLOAD_FOLDER"]
        self.original_static_folder = invoice_app.app.static_folder
        invoice_app.DATABASE_PATH = root / "invoices.sqlite3"
        invoice_app.app.config["UPLOAD_FOLDER"] = str(root / "static" / "uploads")
        invoice_app.app.static_folder = str(root / "static")
        Path(invoice_app.app.config["UPLOAD_FOLDER"]).mkdir(parents=True)
        invoice_app.initialize_database()
        invoice_app.ner_pipeline = None
        invoice_app.ner_status = "unavailable"
        invoice_app.app.config["TESTING"] = True
        self.client = invoice_app.app.test_client()

    def tearDown(self):
        invoice_app.DATABASE_PATH = self.original_database_path
        invoice_app.app.config["UPLOAD_FOLDER"] = self.original_upload_folder
        invoice_app.app.static_folder = self.original_static_folder
        self.temp_dir.cleanup()

    def make_invoice_image(self):
        image = np.full((1050, 1600, 3), 255, dtype=np.uint8)
        lines = [
            "ACME OFFICE SUPPLY",
            "INVOICE NUMBER: INV-2048",
            "INVOICE DATE: 04/12/2025",
            "SUBTOTAL: $120.00",
            "TAX: $9.60",
            "TOTAL AMOUNT: $129.60",
        ]
        for index, line in enumerate(lines):
            cv2.putText(
                image,
                line,
                (100, 130 + index * 135),
                cv2.FONT_HERSHEY_SIMPLEX,
                1.5,
                (15, 15, 15),
                3,
                cv2.LINE_AA,
            )
        ok, encoded = cv2.imencode(".png", image)
        self.assertTrue(ok)
        return encoded.tobytes()

    def test_field_extraction_uses_invoice_labels(self):
        fields = invoice_app.extract_invoice_fields(
            "ACME OFFICE SUPPLY\n"
            "INVOICE NUMBER: INV-2048\n"
            "INVOICE DATE: 04/12/2025\n"
            "SUBTOTAL: $120.00\n"
            "TAX: $9.60\n"
            "TOTAL AMOUNT: $129.60"
        )
        self.assertEqual(fields["Invoice Number"], "INV-2048")
        self.assertEqual(fields["Invoice Date"], "04/12/2025")
        self.assertEqual(fields["Supplier"], "ACME OFFICE SUPPLY")
        self.assertEqual(fields["Subtotal"], "120.00")
        self.assertEqual(fields["Tax"], "9.60")
        self.assertEqual(fields["Total Amount"], "129.60")

    def test_invoice_number_requires_a_labeled_numeric_identifier(self):
        for label in ("Invoice No", "Invoice Number", "Invoice #", "Inv. No."):
            with self.subTest(label=label):
                fields = invoice_app.extract_invoice_fields(
                    f"{label}: ACME-2048\nInvoice Date: 12 Sep 2025"
                )
                self.assertEqual(fields["Invoice Number"], "ACME-2048")

        template_text = (
            "Standard business invoice\n"
            "Your name/business name: [Your logo]\n"
            "Name/business name Invoice #:\n"
            "Business address Invoice date: MM/DD/YYYY\n"
            "Account holder:\n"
            "Subtotal $365.00\nSales tax $35.00"
        )
        fields = invoice_app.extract_invoice_fields(template_text)
        self.assertIsNone(fields["Invoice Number"])
        self.assertIsNone(fields["Invoice Date"])
        self.assertIsNone(fields["Supplier"])
        self.assertIsNone(fields["Total Amount"])
        self.assertEqual(fields["Subtotal"], "365.00")
        self.assertEqual(fields["Tax"], "35.00")

    def test_date_formats_are_calendar_validated_and_due_date_is_not_invoice_date(self):
        cases = {
            "Invoice Date: 2025-09-12": "2025-09-12",
            "Date of Invoice: 09/12/2025": "09/12/2025",
            "Invoice Date: 13/09/2025": "13/09/2025",
            "Issue Date: September 12, 2025": "September 12, 2025",
            "Invoice Date: Sep. 12, 2025": "Sep. 12, 2025",
            "Invoice Date: 31/02/2025": None,
            "Due Date: 09/12/2025": None,
            "Invoice Date: MM/DD/YYYY": None,
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(
                    invoice_app.extract_invoice_fields(text)["Invoice Date"],
                    expected,
                )

    def test_total_prefers_payable_labels_and_ignores_other_numbers(self):
        fields = invoice_app.extract_invoice_fields(
            "Account Number: 987654321\n"
            "PO Number: 456789\n"
            "Subtotal: $1,200.00\n"
            "Tax: $96.00\n"
            "Total before tax: $1,200.00\n"
            "Grand Total: $1,296.00\n"
            "Balance Due: $1,296.00"
        )
        self.assertEqual(fields["Subtotal"], "1200.00")
        self.assertEqual(fields["Tax"], "96.00")
        self.assertEqual(fields["Total Amount"], "1296.00")
        self.assertIsNone(
            invoice_app.extract_invoice_fields(
                "Subtotal: $120.00\nTax: $9.60\nAccount Number: 129.60"
            )["Total Amount"]
        )
        self.assertIsNone(
            invoice_app.extract_invoice_fields(
                "Total:\nAccount Number: 129.60"
            )["Total Amount"]
        )
        self.assertIsNone(
            invoice_app.extract_invoice_fields(
                "Invoice No:\nPO-2025-100"
            )["Invoice Number"]
        )
        self.assertIsNone(
            invoice_app.extract_invoice_fields(
                "Invoice No: 2025-09-12"
            )["Invoice Number"]
        )
        for label in ("Grand Total", "Total Due", "Amount Due", "Balance Due"):
            with self.subTest(label=label):
                fields = invoice_app.extract_invoice_fields(
                    f"Subtotal: $90.00\nTax: $10.00\n{label}: $100.00"
                )
                self.assertEqual(fields["Total Amount"], "100.00")

    def test_supplier_must_be_a_name_not_a_document_label(self):
        template = invoice_app.extract_invoice_fields(
            "Standard business invoice\n"
            "Your name/business name: [Your logo]\n"
            "PO number:\n"
            "Account holder:\n"
            "Swift/BIC:\n"
            "Billed to:\n"
            "Business address Invoice date: MM/DD/YYYY"
        )
        self.assertIsNone(template["Supplier"])

        explicit = invoice_app.extract_invoice_fields(
            "Supplier: Northwind Paper Co.\nInvoice No: NW-5482"
        )
        self.assertEqual(explicit["Supplier"], "Northwind Paper Co.")

    def test_extraction_handles_abbreviated_dates_labels_and_inline_tax_rates(self):
        fields = invoice_app.extract_invoice_fields(
            "Supplier Name: Northwind Paper Co.\n"
            "Invoice Date: Sept. 12, 2025\n"
            "Tax (8%) $9.60\n"
            "Total: $1.00\n"
            "Grand Total: $129.60"
        )
        self.assertEqual(fields["Supplier"], "Northwind Paper Co.")
        self.assertEqual(fields["Invoice Date"], "Sept. 12, 2025")
        self.assertEqual(fields["Tax"], "9.60")
        self.assertEqual(fields["Total Amount"], "129.60")
        business_name = invoice_app.extract_invoice_fields(
            "Business Services International"
        )
        self.assertEqual(business_name["Supplier"], "Business Services International")

    def test_saved_ocr_is_re_extracted_when_extraction_version_changes(self):
        template_text = (
            "Standard business invoice\n"
            "Your name/business name: [Your logo]\n"
            "Name/business name Invoice #:\n"
            "Business address Invoice date: MM/DD/YYYY\n"
            "Subtotal $365.00\nSales tax $35.00"
        )
        old_fields = {
            "Invoice Number": "Business",
            "Invoice Date": None,
            "Supplier": "Your name/business name: [Your logo]",
            "Subtotal": "365.00",
            "Tax": "35.00",
            "Total Amount": None,
        }
        with invoice_app.get_db() as connection:
            cursor = connection.execute(
                """
                INSERT INTO invoices
                    (filename, created_at, status, fields_json, entities_json,
                     extracted_text, image_file, warnings_json)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    "template.png",
                    "2025-01-01T00:00:00+00:00",
                    "review",
                    invoice_app.json.dumps(old_fields),
                    "[]",
                    template_text,
                    "preview.png",
                    "[]",
                ),
            )
            invoice_id = cursor.lastrowid

        invoice_app.upgrade_saved_extractions()
        saved = invoice_app.fetch_invoice(invoice_id)
        self.assertIsNone(saved["fields"]["Invoice Number"])
        self.assertIsNone(saved["fields"]["Invoice Date"])
        self.assertIsNone(saved["fields"]["Supplier"])
        self.assertEqual(saved["fields"]["Subtotal"], "365.00")
        self.assertEqual(saved["fields"]["Tax"], "35.00")
        self.assertIsNone(saved["fields"]["Total Amount"])
        self.assertEqual(saved["field_quality"]["Invoice Number"]["status"], "not_found")

        detail = self.client.get(f"/invoices/{invoice_id}")
        self.assertEqual(detail.status_code, 200)
        detail_html = detail.get_data(as_text=True)
        self.assertRegex(
            detail_html,
            r'<span class="field-label">Invoice Number</span>\s*'
            r"<strong>Not found</strong>",
        )
        self.assertRegex(
            detail_html,
            r'<span class="field-label">Supplier</span>\s*'
            r"<strong>Not found</strong>",
        )
        json_export = self.client.get(f"/export/{invoice_id}.json")
        self.assertEqual(json_export.json["fields"], saved["fields"])
        excel_export = self.client.get(f"/export.xlsx?id={invoice_id}")
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(excel_export.data), read_only=True)
        sheet = workbook["Invoices"]
        headers = next(sheet.iter_rows(values_only=True))
        values = next(sheet.iter_rows(min_row=2, values_only=True))
        for field, value in saved["fields"].items():
            self.assertEqual(values[headers.index(field)], value)
        workbook.close()

    def test_dashboard_and_empty_history_render(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"No invoices yet", response.data)
        self.assertIn(b"InvoiceAI", response.data)
        self.assertIn(b"Extracted data", response.data)
        self.assertEqual(self.client.get("/history").status_code, 200)
        self.assertEqual(self.client.get("/analytics").status_code, 200)
        self.assertEqual(self.client.get("/about").status_code, 200)

    def test_upload_ocr_history_detail_and_exports(self):
        response = self.client.post(
            "/upload",
            data={"invoice": (io.BytesIO(self.make_invoice_image()), "invoice.png")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        detail = self.client.get(response.headers["Location"])
        self.assertEqual(detail.status_code, 200)
        self.assertIn(b"INV-2048", detail.data)
        self.assertIn(b"RULE-BASED EXTRACTION", detail.data)
        self.assertIn(b"TRANSFORMER NER", detail.data)
        self.assertIn(b"OCR OUTPUT", detail.data)
        self.assertIn(b"Highlighted invoice", detail.data)
        self.assertIn(b"TOTAL AMOUNT", detail.data)
        self.assertIn(b"Label matched", detail.data)

        history = self.client.get("/history?q=INV-2048")
        self.assertEqual(history.status_code, 200)
        self.assertIn(b"invoice.png", history.data)
        self.assertIn(b"129.60", history.data)
        filtered = self.client.get(
            "/history?sort=total_amount&direction=asc"
            "&date_from=2020-01-01&date_to=2030-01-01"
        )
        self.assertEqual(filtered.status_code, 200)
        self.assertIn(b"invoice.png", filtered.data)

        invoice_id = int(response.headers["Location"].rsplit("/", 1)[1])
        saved_invoice = invoice_app.fetch_invoice(invoice_id)
        preview_response = self.client.get(
            f"/static/uploads/{saved_invoice['image_file']}"
        )
        self.assertEqual(preview_response.status_code, 200)
        self.assertTrue(preview_response.data.startswith(b"\x89PNG"))
        preview_image = cv2.imdecode(
            np.frombuffer(preview_response.data, dtype=np.uint8),
            cv2.IMREAD_COLOR,
        )
        self.assertEqual(preview_image.shape[:2], (1050, 1600))
        preview_response.close()
        exported_json = self.client.get(f"/export/{invoice_id}.json")
        self.assertEqual(exported_json.status_code, 200)
        self.assertEqual(exported_json.json["fields"], saved_invoice["fields"])
        self.assertEqual(
            exported_json.json["fields"]["Invoice Number"], "INV-2048"
        )
        self.assertEqual(
            exported_json.json["field_quality"]["Invoice Number"]["status"],
            "matched",
        )
        self.assertIn("extracted_text", exported_json.json)

        exported_excel = self.client.get(f"/export.xlsx?id={invoice_id}")
        self.assertEqual(exported_excel.status_code, 200)
        self.assertTrue(exported_excel.data.startswith(b"PK"))
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(exported_excel.data), read_only=True)
        invoice_sheet = workbook["Invoices"]
        headers = next(invoice_sheet.iter_rows(values_only=True))
        values = next(invoice_sheet.iter_rows(min_row=2, values_only=True))
        for field, value in saved_invoice["fields"].items():
            self.assertEqual(values[headers.index(field)], value)
        quality_rows = list(
            workbook["Field quality"].iter_rows(values_only=True)
        )
        self.assertIn(("invoice.png", "Invoice Number", "matched",
                       "A numeric identifier follows an invoice-number label."),
                      quality_rows)
        workbook.close()
        self.assertEqual(self.client.get("/export.xlsx").status_code, 200)

    def test_pdf_upload_processes_first_page(self):
        pdf_data = io.BytesIO()
        Image.open(io.BytesIO(self.make_invoice_image())).save(
            pdf_data, format="PDF", resolution=150
        )
        response = self.client.post(
            "/upload",
            data={"invoice": (io.BytesIO(pdf_data.getvalue()), "invoice.pdf")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        detail = self.client.get(response.headers["Location"])
        self.assertEqual(detail.status_code, 200)
        self.assertIn(b"INV-2048", detail.data)

    def test_upload_size_limit_returns_friendly_error(self):
        original_limit = invoice_app.app.config["MAX_CONTENT_LENGTH"]
        invoice_app.app.config["MAX_CONTENT_LENGTH"] = 128
        try:
            response = self.client.post(
                "/upload",
                data={"invoice": (io.BytesIO(b"x" * 256), "invoice.png")},
                content_type="multipart/form-data",
            )
        finally:
            invoice_app.app.config["MAX_CONTENT_LENGTH"] = original_limit
        self.assertEqual(response.status_code, 413)
        self.assertIn(b"exceeds the 128 bytes upload limit", response.data)

    def test_empty_unsupported_and_unreadable_uploads_are_rejected(self):
        empty = self.client.post(
            "/upload",
            data={"invoice": (io.BytesIO(b""), "empty.png")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertEqual(empty.status_code, 200)
        self.assertIn(b"file is empty", empty.data)

        unsupported = self.client.post(
            "/upload",
            data={"invoice": (io.BytesIO(b"not an invoice"), "notes.txt")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertEqual(unsupported.status_code, 200)
        self.assertIn(b"Unsupported format", unsupported.data)

        unreadable = self.client.post(
            "/upload",
            data={"invoice": (io.BytesIO(b"not an image"), "broken.png")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertEqual(unreadable.status_code, 200)
        self.assertIn(b"could not be read", unreadable.data)

    def test_unknown_invoice_returns_friendly_not_found(self):
        response = self.client.get("/invoices/9999")
        self.assertEqual(response.status_code, 404)
        self.assertIn(b"could not be found", response.data)


if __name__ == "__main__":
    unittest.main()
